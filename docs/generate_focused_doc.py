# -*- coding: utf-8 -*-
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="3182CE"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(45, 55, 72)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_callout(doc, title, text, bg_color="EBF8FF", border_color="3182CE"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"📌 {title}: ")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(43, 108, 176)
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(45, 55, 72)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def add_when_box(doc, when_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "FEFCBF") # Light amber
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="D69E2E"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    rt = p.add_run("⏰ KHI NÀO DÙNG API NÀY?\n")
    rt.font.bold = True
    rt.font.size = Pt(10.5)
    rt.font.color.rgb = RGBColor(146, 64, 14) # Dark Amber
    
    rd = p.add_run(when_text)
    rd.font.size = Pt(9.5)
    rd.font.color.rgb = RGBColor(45, 55, 72)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def main():
    doc = docx.Document()
    
    # 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.different_first_page_header_footer = True
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("SBA301 - Hướng Dẫn Sử Dụng 5 Nhóm API Trọng Tâm (Function 4)")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(160, 174, 192)

    # ------------------ TITLE & COVER ------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("HƯỚNG DẪN SỬ DỤNG 5 NHÓM API TRỌNG TÂM\nVÀ THỜI ĐIỂM KÍCH HOẠT TRONG QUY TRÌNH THUÊ XE")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(26, 54, 93)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Tập trung 5 nghiệp vụ: Deposit (Cọc), Extra Fees (Phụ phí), Disputes (Khiếu nại), Payouts (Chi trả), Review (Đánh giá)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(74, 85, 104)

    # Meta box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_table, color="E2E8F0", sz="6")
    meta_data = [
        ("Dự án áp dụng:", "SBA301 - Nền tảng Cho Thuê Xe Tự Lái (Car Rental Marketplace)"),
        ("Phân hệ chức năng:", "Function 4 - Quản lý Tài chính, Phí phát sinh, Khiếu nại & Đánh giá chuyến đi"),
        ("Nhóm API trọng tâm:", "1. Deposit | 2. Extra-Fees | 3. Disputes | 4. Payouts | 5. Reviews"),
        ("Mục tiêu tài liệu:", "Chỉ rõ chức năng, endpoint API và thời điểm chính xác khi nào kích hoạt từng API trong thực tế")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        c0 = meta_table.cell(row_idx, 0)
        c1 = meta_table.cell(row_idx, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F7FAFC")
        set_cell_margins(c0, 50, 50, 90, 90)
        set_cell_margins(c1, 50, 50, 90, 90)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(43, 108, 176)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(45, 55, 72)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ------------------ BẢNG TỔNG QUAN ------------------
    h_overview = doc.add_heading("TỔNG QUAN: BẢNG TRA CỨU THỜI ĐIỂM SỬ DỤNG 5 NHÓM API", level=1)
    h_overview.paragraph_format.space_before = Pt(12)
    h_overview.paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Bảng tra cứu nhanh giúp lập trình viên Frontend, Backend và Tester nắm bắt thời điểm gọi từng API:").paragraph_format.space_after = Pt(4)

    sum_table = doc.add_table(rows=1, cols=5)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sum_table, color="CBD5E0", sz="4")

    sum_headers = ["Nhóm API", "Endpoint chính", "Người gọi", "Thời điểm kích hoạt trong quy trình", "Mục đích"]
    sum_widths = [Inches(1.1), Inches(1.8), Inches(0.9), Inches(1.8), Inches(1.5)]
    
    for i, title in enumerate(sum_headers):
        sum_table.rows[0].cells[i].width = sum_widths[i]
        set_cell_background(sum_table.rows[0].cells[i], "2B6CB0")
        set_cell_margins(sum_table.rows[0].cells[i], 80, 80, 70, 70)
        p = sum_table.rows[0].cells[i].paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    quick_rows = [
        ("1. DEPOSIT\n(Đặt cọc)", "/api/transactions/deposit\n/refund", "Khách thuê /\nAdmin", "Ngay sau khi Chủ xe duyệt đơn (CONFIRMED) và trước giờ nhận xe", "Giữ xe an toàn trên sàn; hoàn lại khi kết thúc chuyến xe nguyên vẹn"),
        ("2. EXTRA-FEES\n(Phí phát sinh)", "/api/extra-fees\n/pay | /waive", "Chủ xe /\nKhách thuê", "Lúc trả xe (Return Handover) phát hiện vi phạm (bẩn xe, trễ giờ, quá km)", "Thu thêm phụ phí hoặc miễn trừ nếu chủ xe thông cảm"),
        ("3. DISPUTES\n(Khiếu nại)", "/api/disputes\n/evidences | /resolve", "Khách / Chủ xe\n(Admin xử)", "Khi phát sinh bất đồng về tiền cọc, phụ phí hoặc tình trạng xe bàn giao", "Trọng tài Admin phân xử công bằng, bảo vệ tiền ký quỹ của 2 bên"),
        ("4. PAYOUTS\n(Chi trả chủ xe)", "/api/payouts\n/process", "Hệ thống sàn /\nKế toán", "Sau khi chuyến đi hoàn tất (COMPLETED) và không còn khiếu nại tranh chấp", "Quyết toán tiền thuê (trừ 10% phí sàn) chuyển về tài khoản ngân hàng chủ xe"),
        ("5. REVIEWS\n(Đánh giá)", "/api/reviews", "Khách thuê", "Sau khi trả xe hoàn tất (COMPLETED) và kết thúc chuyến đi", "Chấm điểm 1-5 sao và nhận xét thực tế về xe và chủ xe")
    ]

    for col0, col1, col2, col3, col4 in quick_rows:
        row = sum_table.add_row().cells
        for idx, text in enumerate([col0, col1, col2, col3, col4]):
            row[idx].width = sum_widths[idx]
            set_cell_margins(row[idx], 60, 60, 60, 60)
            p = row[idx].paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(26, 54, 93)
            elif idx == 1:
                r.font.name = "Consolas"
                r.font.color.rgb = RGBColor(43, 108, 176)
            elif idx == 3:
                r.bold = True
                r.font.color.rgb = RGBColor(146, 64, 14)
            else:
                r.font.color.rgb = RGBColor(45, 55, 72)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ------------------ CHI TIẾT 1. DEPOSIT ------------------
    h1 = doc.add_heading("1. NHÓM API ĐẶT CỌC & HOÀN CỌC (DEPOSIT & REFUND)", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)

    add_when_box(
        doc,
        "1. KHI NÀO TẠO ĐẶT CỌC (Deposit)?\n"
        "   - Sau khi chủ xe đã bấm xác nhận duyệt đơn thuê (OrderStatus = CONFIRMED).\n"
        "   - Trước thời hạn thanh toán (expiresAt - thường là trong vòng 2 tiếng sau khi duyệt).\n"
        "   - Khách thuê bắt buộc phải đặt cọc số tiền quy định (ví dụ 2,000,000 VND) để giữ xe chắc chắn.\n\n"
        "2. KHI NÀO DÙNG HOÀN TIỀN CỌC (Refund)?\n"
        "   - Trường hợp 1 (Bình thường): Khách trả xe nguyên vẹn, biên bản trả xe đã ký xong (OrderStatus = COMPLETED). Admin gọi API hoàn lại 100% tiền cọc cho khách.\n"
        "   - Trường hợp 2 (Hủy đơn): Khách hoặc chủ xe hủy đơn trước ngày đi (hủy hợp lệ theo chính sách), Admin gọi hoàn tiền cọc lại cho khách.\n"
        "   - Trường hợp 3 (Giải quyết khiếu nại): Admin phân xử khiếu nại xong, kích hoạt hoàn cọc theo tỷ lệ phán quyết."
    )

    doc.add_paragraph("Các API thành phần trong nhóm Deposit:").paragraph_format.space_after = Pt(2)

    doc.add_paragraph("a) Tạo giao dịch Đặt cọc:", style='List Bullet').runs[0].bold = True
    doc.add_paragraph("   - Method & URL: POST /api/transactions/deposit\n   - Header: Authorization: Bearer <TOKEN_KHACH>\n   - Body:")
    add_code_block(doc, """// POST /api/transactions/deposit
{
  "orderId": 1,
  "amount": 2000000,
  "method": "VNPAY" // hoặc "MOCK_GATEWAY"
}

// Response (HTTP 200):
{
  "success": true,
  "message": "Tạo giao dịch đặt cọc thành công",
  "data": {
    "id": 102,
    "orderId": 1,
    "transactionType": "DEPOSIT",
    "amount": 2000000,
    "status": "PENDING",
    "paymentUrl": "https://sandbox.vnpayment.vn/paymentv2/vpcpay.html?..."
  }
}""")

    doc.add_paragraph("b) Hoàn tiền cọc sau chuyến đi:", style='List Bullet').runs[0].bold = True
    doc.add_paragraph("   - Method & URL: POST /api/transactions/refund\n   - Header: Authorization: Bearer <TOKEN_ADMIN>\n   - Body:")
    add_code_block(doc, """// POST /api/transactions/refund
{
  "parentTransactionId": 102, // ID giao dịch cọc ban đầu (đã SUCCESS)
  "amount": 2000000,          // Số tiền cọc cần hoàn trả
  "reason": "Hoàn tiền cọc sau chuyến đi khi kiểm tra xe nguyên vẹn"
}""")

    # ------------------ CHI TIẾT 2. EXTRA FEES ------------------
    h2 = doc.add_heading("2. NHÓM API PHÍ PHÁT SINH (EXTRA-FEES)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)

    add_when_box(
        doc,
        "1. KHI NÀO TẠO PHÍ PHÁT SINH (POST /api/extra-fees)?\n"
        "   - Diễn ra vào thời điểm trả xe (Return Handover) khi kết thúc chuyến đi.\n"
        "   - Khi chủ xe kiểm tra hiện trạng xe và phát hiện các trường hợp:\n"
        "     * Khách trả xe quá bẩn (nội thất đổ đồ ăn, bùn đất) -> Thu phí vệ sinh (CLEANING, ví dụ 150,000 VND).\n"
        "     * Khách trả trễ hẹn so với giờ đăng ký -> Thu phí trễ giờ (LATE_RETURN, ví dụ 100,000 VND/giờ).\n"
        "     * Số km thực tế vượt hạn mức quy định (vd > 300km/ngày) -> Thu phí quá km (EXTRA_KM, 5,000 VND/km).\n"
        "     * Xe có vết xước nhỏ cần đánh bóng -> Thu phí khắc phục (DAMAGE).\n\n"
        "2. KHI NÀO THANH TOÁN PHỤ PHÍ (POST /api/extra-fees/{id}/pay)?\n"
        "   - Sau khi chủ xe tạo phụ phí, khách thuê đồng ý với chi phí này và tiến hành thanh toán trực tuyến qua sàn.\n\n"
        "3. KHI NÀO MIỄN TRỪ PHỤ PHÍ (PUT /api/extra-fees/{id}/waive)?\n"
        "   - Khi chủ xe hoặc Admin quyết định thông cảm bỏ qua (ví dụ: khách thân thiết, khách trễ ít phút do kẹt xe hoặc khách tự rửa xe ngoài sạch sẽ rồi)."
    )

    doc.add_paragraph("Các API thành phần trong nhóm Extra-Fees:").paragraph_format.space_after = Pt(2)

    doc.add_paragraph("a) Chủ xe tạo phí phát sinh:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// POST /api/extra-fees
{
  "orderId": 1,
  "feeType": "CLEANING", // CLEANING, LATE_RETURN, EXTRA_KM, DAMAGE, OTHER
  "amount": 150000,
  "description": "Xe bị dính bùn đất nội thất và mùi hải sản, cần vệ sinh khử mùi"
}

// Response: HTTP 200 OK -> ExtraFee id=43, status=PENDING""")

    doc.add_paragraph("b) Thanh toán phí phát sinh (Hỗ trợ cả 2 phương thức):", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// POST /api/extra-fees/43/pay

// PHƯƠNG ÁN 1: Thanh toán qua VNPay (Khách thuê thanh toán trực tuyến)
// Request Body:
{
  "method": "VNPAY"
}
// Quyền gọi: Người thuê xe (Renter)
// Kết quả: Tạo giao dịch PENDING và trả link VNPay (paymentUrl) để khách thanh toán.
// Sau khi thanh toán thành công qua VNPay IPN, phụ phí tự động chuyển sang PAID.

// PHƯƠNG ÁN 2: Khách trả tiền mặt trực tiếp cho chủ xe lúc trả xe (CASH / DIRECT)
// Request Body:
{
  "method": "CASH" // hoặc "DIRECT"
}
// Quyền gọi: Chủ xe (hoặc Admin) bấm xác nhận đã nhận đủ tiền mặt từ khách
// Kết quả: Phụ phí chuyển ngay sang PAID và tạo Transaction SUCCESS loại EXTRA_FEE.""")

    doc.add_paragraph("c) Miễn trừ phụ phí:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// PUT /api/extra-fees/43/waive?note=Khach+hang+than+thiet+nen+mien+giam
// Chuyển trạng thái ExtraFee sang WAIVED (không cần trả tiền)""")

    # ------------------ CHI TIẾT 3. DISPUTES ------------------
    h3 = doc.add_heading("3. NHÓM API KHIẾU NẠI & TRANH CHẤP (DISPUTES)", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)

    add_when_box(
        doc,
        "1. KHI NÀO MỞ KHIẾU NẠI (POST /api/disputes)?\n"
        "   - Khi giữa Khách thuê và Chủ xe không đạt được thỏa thuận chung về các vấn đề tài chính hoặc quyền lợi:\n"
        "     * Tình huống 1: Khách không đồng ý với khoản phí phát sinh chủ xe đòi thu (ví dụ chủ xe báo trầy xước nhưng vết xước đó đã có từ trước chuyến đi).\n"
        "     * Tình huống 2: Khách trả xe đã lâu nhưng chủ xe chần chừ không xác nhận biên bản trả xe để hệ thống hoàn cọc.\n"
        "     * Tình huống 3: Xe gặp sự cố hỏng hóc kỹ thuật giữa đường do lỗi xe cũ của chủ xe.\n\n"
        "2. KHI NÀO BỔ SUNG BẰNG CHỨNG (POST /api/disputes/{id}/evidences)?\n"
        "   - Ngay sau khi mở khiếu nại, một trong hai bên cung cấp hình ảnh/video bằng chứng (ảnh biên bản bàn giao lúc nhận xe, ảnh chụp đồng hồ ODO, video quay xe trước chuyến đi) để chứng minh mình đúng.\n\n"
        "3. KHI NÀO ADMIN PHÂN XỬ (PUT /api/disputes/{id}/resolve)?\n"
        "   - Khi Admin xem xét toàn bộ bằng chứng từ 2 bên và đưa ra phán quyết cuối cùng.\n"
        "   - TỰ ĐỘNG HOÀN TIỀN: Nếu phán quyết có số tiền bồi hoàn (approvedAmount > 0), hệ thống tự động kích hoạt tạo giao dịch REFUND hoàn tiền cho khách ngay lập tức!"
    )

    doc.add_paragraph("Các API thành phần trong nhóm Disputes:").paragraph_format.space_after = Pt(2)

    doc.add_paragraph("a) Khách hoặc Chủ xe mở khiếu nại:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// POST /api/disputes
{
  "orderId": 1,
  "extraFeeId": 43, // Tùy chọn (nếu khiếu nại về phụ phí cụ thể)
  "reason": "Không đồng ý với phí vệ sinh nội thất vì khi nhận xe sàn đã dơ sẵn"
}

// Response: HTTP 200 OK -> Dispute id=26, status=OPEN""")

    doc.add_paragraph("b) Đăng tải bằng chứng (ảnh/video):", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// POST /api/disputes/26/evidences
{
  "fileUrl": "https://images.unsplash.com/photo-1549317661-bd32c8ce0db2",
  "description": "Ảnh chụp sàn xe lúc nhận xe đã có vết bẩn lưu trong biên bản Pickup"
}""")

    doc.add_paragraph("c) Admin ra quyết định phân xử & Tự động hoàn tiền:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// PUT /api/disputes/26/resolve
{
  "status": "RESOLVED",
  "resolutionNote": "Admin xác nhận vết bẩn đã có từ trước, chấp thuận bồi hoàn 150,000 VND cho khách",
  "approvedAmount": 150000 // Số tiền phê duyệt hoàn lại cho khách
}

// CƠ CHẾ TỰ ĐỘNG THỰC THI PHÁN QUYẾT:
// 1. Chuyển trạng thái Dispute sang RESOLVED.
// 2. Tự động kích hoạt gọi giao dịch REFUND hoàn tiền approvedAmount cho khách.
// 3. Nếu khiếu nại phụ phí đang PENDING -> Tự động chuyển phụ phí sang WAIVED.""")

    # ------------------ CHI TIẾT 4. PAYOUTS ------------------
    h4 = doc.add_heading("4. NHÓM API CHI TRẢ DOANH THU CHỦ XE (PAYOUTS)", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)

    add_when_box(
        doc,
        "1. KHI NÀO TẠO LỆNH CHI TRẢ (POST /api/payouts)?\n"
        "   - Sau khi chuyến đi đã kết thúc hoàn tất (bắt buộc OrderStatus = COMPLETED).\n"
        "   - Kiểm tra chống trùng lặp: Mỗi đơn hàng chỉ được tạo chi trả 1 lần duy nhất.\n"
        "   - Tự động hóa tính toán: Hệ thống tự trích xuất thông tin chủ xe và tính toán các khoản tiền từ đơn hàng:\n"
        "     * grossAmount: Tổng giá trị đơn hàng (order.totalAmount)\n"
        "     * platformFee: Phí hoa hồng sàn theo tỷ lệ (mặc định 10%)\n"
        "     * netAmount = grossAmount - platformFee (số tiền thực nhận của chủ xe)\n\n"
        "2. KHI NÀO XỬ LÝ LỆNH CHI TRẢ (PUT /api/payouts/{id}/process)?\n"
        "   - Kế toán sàn hoặc hệ thống thanh toán tự động thực hiện chuyển khoản ngân hàng tới số tài khoản đã liên kết của chủ xe (bankAccountId).\n"
        "   - Sau khi tiền đã vào tài khoản chủ xe thành công, gọi API này để chuyển trạng thái Payout sang PAID."
    )

    doc.add_paragraph("Các API thành phần trong nhóm Payouts:").paragraph_format.space_after = Pt(2)

    doc.add_paragraph("a) Tạo lệnh chi trả doanh thu:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// POST /api/payouts
// Quyền gọi: Admin
{
  "orderId": 1,
  "note": "Chi trả tiền thuê đơn ORD-2026-001 (khấu trừ 10% phí hoa hồng sàn)"
}

// Hệ thống tự động xác định chủ xe và tính toán:
// - grossAmount = 1,750,000 VND
// - platformFee = 175,000 VND (10%)
// - netAmount = 1,575,000 VND
// Response: HTTP 200 OK -> Payout id=25, netAmount=1575000, status=PENDING""")

    doc.add_paragraph("b) Cập nhật trạng thái sau khi chuyển khoản thành công:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// PUT /api/payouts/25/process
// Chuyển Payout sang PAID -> Hoàn tất chu trình tài chính cho chuyến đi!""")

    # ------------------ CHI TIẾT 5. REVIEWS ------------------
    h5 = doc.add_heading("5. NHÓM API ĐÁNH GIÁ CHUYẾN ĐI (REVIEWS)", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(4)

    add_when_box(
        doc,
        "1. KHI NÀO DÙNG API ĐÁNH GIÁ (POST /api/reviews)?\n"
        "   - Kích hoạt khi chuyến đi đã hoàn tất (OrderStatus = COMPLETED).\n"
        "   - Khách thuê đã trả xe và nhận lại tiền cọc.\n"
        "   - Ứng dụng hiển thị màn hình 'Đánh giá chuyến đi' để khách cho điểm sao (từ 1 đến 5 sao) và nhận xét về:\n"
        "     * Độ sạch sẽ, máy móc của xe, trang thiết bị đúng như mô tả.\n"
        "     * Thái độ phục vụ, tính đúng giờ và sự hỗ trợ của chủ xe.\n\n"
        "2. RÀNG BUỘC NGHIỆP VỤ QUAN TRỌNG:\n"
        "   - Mỗi đơn thuê xe (orderId) chỉ được phép gửi đánh giá 1 lần duy nhất (ngăn chặn spam đánh giá ảo).\n"
        "   - Nếu khách đã đánh giá rồi mà gửi lại, backend sẽ trả về mã lỗi 400 Bad Request: 'Bạn đã gửi đánh giá cho đơn thuê này rồi'."
    )

    doc.add_paragraph("API thành phần trong nhóm Reviews:").paragraph_format.space_after = Pt(2)

    doc.add_paragraph("Khách gửi đánh giá sau chuyến đi:", style='List Bullet').runs[0].bold = True
    add_code_block(doc, """// POST /api/reviews
{
  "orderId": 1,
  "reviewerId": 1, // Khách thuê
  "revieweeId": 2, // Chủ xe
  "carId": 3,
  "rating": 5,     // Thang điểm 1 - 5 sao
  "comment": "Xe chạy rất êm, sạch sẽ thơm tho, chủ xe nhiệt tình hỗ trợ nhận xe sớm 15 phút. 5 sao!"
}

// Response: HTTP 200 OK -> Review id=1 được lưu và tự tính vào điểm trung bình của xe/chủ xe.""")

    # ------------------ LỆNH KIỂM THỬ POSTMAN ------------------
    h_test = doc.add_heading("HƯỚNG DẪN KIỂM THỬ 5 NHÓM API BẰNG POSTMAN & NEWMAN", level=1)
    h_test.paragraph_format.space_before = Pt(14)
    h_test.paragraph_format.space_after = Pt(4)

    doc.add_paragraph(
        "Cả 5 nhóm API trên đã được đóng gói sẵn trong file Postman Collection tại thư mục docs/carrental_postman_collection.json. "
        "Người dùng có thể kiểm thử tự động toàn bộ 5 nhóm nghiệp vụ này bằng lệnh Newman CLI:"
    ).paragraph_format.line_spacing = 1.15

    add_code_block(doc, """# Di chuyển vào thư mục backend và chạy kiểm thử tự động:
cd d:\\Study\\SBA302\\project\\SBA301\\carrental-BE
npx --yes newman run carrental_postman_collection.json

# Tài khoản đăng nhập test:
# Email: admin@carrental.com
# Mật khẩu: 123456
# Kết quả kiểm thử: 26/26 Requests executed, 40/40 Assertions Passed 100%""")

    target_path = "d:/Study/SBA302/project/SBA301/docs/HUONG_DAN_SU_DUNG_API_FUNCTION_4.docx"
    doc.save(target_path)
    print("Created:", target_path)

if __name__ == "__main__":
    main()
